#!/usr/bin/env python3
"""ThirdC 文档解析器：PDF / DOCX → 纯文本（供 Rust 内核调用）"""
import sys, json

def extract_pdf(path):
    import fitz
    doc = fitz.open(path)
    pages = []
    for page in doc:
        pages.append(page.get_text())
    return "\n\n".join(pages)

def extract_docx(path):
    from docx import Document
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t: continue
        style = para.style.name.lower() if para.style else ""
        if "heading" in style:
            level = "##"
            try: n = int(style.replace("heading","").strip())
            except: n = 2
            level = "#" * max(2, min(n, 4))
            parts.append(f"{level} {t}")
        else:
            parts.append(t)
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("|","\\|") for c in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)

if __name__ == "__main__":
    path = sys.argv[1]
    ext = path.rsplit(".",1)[-1].lower()
    try:
        if ext == "pdf": text = extract_pdf(path)
        elif ext == "docx": text = extract_docx(path)
        else: text = open(path, encoding="utf-8", errors="replace").read()
        print(json.dumps({"ok": True, "text": text}))
    except Exception as e:
        print(json.dumps({"ok": False, "error": str(e)}))
